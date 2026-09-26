import csv
import io
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.analytics import insights as insight_engine
from apps.analytics import reports
from apps.analytics.models import InsightReport
from apps.delivery.models import DeliverySettings, FeeSlab
from apps.delivery.services import compute_fee, delivery_cost

from ..permissions import dash
from .orders import _safe

RANGES = [("7", "Last 7 days"), ("30", "Last 30 days"), ("90", "Last 90 days"), ("365", "Last 12 months")]
REPORTS = {
    "daily": (
        "Sales by day",
        ["Day", "Orders", "Sales"],
        lambda p: [[r["day"], r["orders"], r["revenue"]] for r in reports.daily(p)],
    ),
    "products": (
        "Top products",
        ["Product", "Size", "Units", "Sales"],
        lambda p: [
            [r["variant__product__name"], r["variant_label"], r["qty"], r["revenue"]]
            for r in reports.top_products(p, 200)
        ],
    ),
    "categories": (
        "Sales by category",
        ["Category", "Units", "Sales"],
        lambda p: [[r["category"], r["qty"], r["revenue"]] for r in reports.by_category(p)],
    ),
    "areas": (
        "Sales by area",
        ["Pincode", "Area", "Orders", "Sales", "Delivery fees", "Delivery cost"],
        lambda p: [
            [r["ship_pincode"], r["ship_area"], r["orders"], r["revenue"], r["fees"], r["cost"]]
            for r in reports.by_area(p, 200)
        ],
    ),
    "payments": (
        "Payment methods",
        ["Method", "Orders", "Sales"],
        lambda p: [[r["payment_method"], r["orders"], r["revenue"]] for r in reports.by_payment(p)],
    ),
    "coupons": (
        "Coupons",
        ["Code", "Uses", "Discount given", "Order sales"],
        lambda p: [[r["coupon__code"], r["uses"], r["discount"], r["sales"]] for r in reports.coupons(p)],
    ),
}


def _period(request):
    try:
        start = date.fromisoformat(request.GET.get("from", ""))
        end = date.fromisoformat(request.GET.get("to", ""))
        if start <= end and (end - start).days <= 731:
            return reports.Period.between(start, end), "custom"
    except ValueError:
        pass
    days = request.GET.get("range", "30")
    days = days if days in dict(RANGES) else "30"
    return reports.Period.last_days(int(days)), days


@dash("manager")
def analytics(request):
    p, rng = _period(request)
    export = request.GET.get("export")
    if export:
        return _export(p, request.GET.get("report", "daily"), export)
    daily = reports.daily(p)
    cats = reports.by_category(p)[:8]
    top = reports.top_products(p, 10)
    grid, peak = reports.heatmap(p)
    chart = {
        "daily": {
            "labels": [d["day"].strftime("%d %b") for d in daily],
            "revenue": [float(d["revenue"]) for d in daily],
            "orders": [d["orders"] for d in daily],
        },
        "top": {
            "labels": [f"{t['variant__product__name']} {t['variant_label']}" for t in top],
            "values": [float(t["revenue"]) for t in top],
        },
        "cats": {"labels": [c["category"] for c in cats], "values": [float(c["revenue"]) for c in cats]},
    }
    return render(
        request,
        "dashboard/analytics.html",
        {
            "k": reports.compare(p),
            "p": p,
            "rng": rng,
            "ranges": RANGES,
            "chart": chart,
            "areas": reports.by_area(p),
            "payments": reports.by_payment(p),
            "coupons": reports.coupons(p),
            "funnel": reports.funnel(p),
            "idle": reports.not_selling(p),
            "risks": reports.stock_risks(p),
            "heat": [
                (
                    ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"][i],
                    [(n, round(n / peak, 2) if peak else 0) for n in row],
                )
                for i, row in enumerate(grid)
            ],
            "reports": [(k, v[0]) for k, v in REPORTS.items()],
            "qs": request.GET.urlencode(),
        },
    )


def _export(p, key, fmt):
    title, headers, fn = REPORTS.get(key, REPORTS["daily"])
    rows = fn(p)
    stamp = f"{timezone.localtime(p.start):%Y%m%d}-{timezone.localtime(p.end - timedelta(seconds=1)):%Y%m%d}"
    name = f"rasiko-{key}-{stamp}"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = title[:30]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="3E2723")
            c.fill = PatternFill("solid", fgColor="FFF4E0")
        for r in rows:
            ws.append([_safe(x) if isinstance(x, str) else (float(x) if isinstance(x, Decimal) else x) for x in r])
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(
            buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
        return resp
    if fmt == "pdf":
        resp = HttpResponse(_pdf(title, headers, rows, p), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
        return resp
    resp = HttpResponse(content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
    resp.write("﻿")
    w = csv.writer(resp)
    w.writerow(headers)
    for r in rows:
        w.writerow([_safe(x) for x in r])
    return resp


def _pdf(title, headers, rows, p):
    # Only escapes text for the PDF; nothing is parsed.
    from xml.sax.saxutils import escape  # nosec B406

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.orders.invoice import _fonts

    _fonts()
    regular, bold = "Dejavu", "Dejavu-Bold"
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=36, rightMargin=36)
    st = getSampleStyleSheet()
    st["Title"].fontName = bold
    st["Normal"].fontName = regular
    period = f"{timezone.localtime(p.start):%d %b %Y} to {timezone.localtime(p.end - timedelta(seconds=1)):%d %b %Y}"
    data = [headers] + [[escape(str(x)) if not isinstance(x, Decimal) else f"₹{x:,.2f}" for x in r] for r in rows]
    t = Table(data, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), regular),
                ("FONTNAME", (0, 0), (-1, 0), bold),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFF4E0")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#3E2723")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#F3E3CC")),
            ]
        )
    )
    doc.build([Paragraph(f"Rasiko · {escape(title)}", st["Title"]), Paragraph(period, st["Normal"]), Spacer(1, 12), t])
    return buf.getvalue()


# ---- Delivery economics and what-if ------------------------------------------------------------------------------

WHATIF_FIELDS = [
    ("free_near_threshold", "Free delivery above (nearby) ₹"),
    ("free_near_max_km", "“Nearby” means within km"),
    ("free_anywhere_threshold", "Free delivery anywhere above ₹"),
    ("small_order_below", "Small order fee below ₹"),
    ("small_order_fee", "Small order fee ₹"),
    ("min_order_value", "Minimum order ₹"),
]


@dash("manager")
def delivery_economics(request):
    cfg = DeliverySettings.load()
    slabs = list(FeeSlab.objects.order_by("min_km"))
    days = 30
    p = reports.Period.last_days(days)
    orders = list(
        reports.orders_in(p)
        .filter(status__in=reports.COUNTED)
        .values("subtotal", "discount", "coins_value", "road_km", "delivery_fee", "small_order_fee", "delivery_cost")
    )
    trial = DeliverySettings(**{f.attname: getattr(cfg, f.attname) for f in cfg._meta.concrete_fields})
    slab_fees = {s.pk: s.fee for s in slabs}
    changed = False
    for key, _label in WHATIF_FIELDS:
        raw = request.GET.get(key, "").strip()
        if raw:
            try:
                setattr(trial, key, Decimal(raw))
                changed = True
            except InvalidOperation:
                messages.error(request, f"{_label}: enter a number.")
    trial_slabs = []
    for s in slabs:
        raw = request.GET.get(f"slab_{s.pk}", "").strip()
        s2 = FeeSlab(min_km=s.min_km, max_km=s.max_km, fee=s.fee)
        if raw:
            try:
                s2.fee = Decimal(raw)
                changed = changed or s2.fee != slab_fees[s.pk]
            except InvalidOperation:
                pass
        trial_slabs.append(s2)
    noon = timezone.localtime().replace(hour=13, minute=0)
    actual_fees = sum((o["delivery_fee"] + o["small_order_fee"] for o in orders), Decimal("0"))
    cost = sum((o["delivery_cost"] or delivery_cost(o["road_km"], cfg) for o in orders), Decimal("0"))
    trial_fees, lost = Decimal("0"), 0
    for o in orders:
        basket = o["subtotal"] - o["discount"] - o["coins_value"]
        q = compute_fee(basket, o["road_km"], trial, trial_slabs, now=noon)
        if q.ok:
            trial_fees += q.delivery_fee + q.small_order_fee
        else:
            lost += 1
    per_slab = []
    for s in slabs:
        mid = (s.min_km + s.max_km) / 2
        per_slab.append({"slab": s, "cost": delivery_cost(Decimal(mid), cfg)})
    return render(
        request,
        "dashboard/delivery_economics.html",
        {
            "cfg": cfg,
            "slabs": slabs,
            "per_slab": per_slab,
            "orders": len(orders),
            "days": days,
            "actual_fees": actual_fees,
            "cost": cost,
            "net": actual_fees - cost,
            "trial_fees": trial_fees,
            "trial_net": trial_fees - cost,
            "lost": lost,
            "changed": changed,
            "fields": [(k, label, request.GET.get(k, ""), getattr(cfg, k)) for k, label in WHATIF_FIELDS],
            "slab_inputs": [(s, request.GET.get(f"slab_{s.pk}", "")) for s in slabs],
        },
    )


# ---- Insights -------------------------------------------------------------------------------------------------


@dash("manager")
def insights(request):
    return render(
        request,
        "dashboard/insights.html",
        {"report": InsightReport.objects.first(), "history": InsightReport.objects.all()[1:8]},
    )


@dash("manager")
@require_POST
def insights_refresh(request):
    insight_engine.build(int(request.POST.get("days", 30)) if request.POST.get("days") in ("7", "30", "90") else 30)
    messages.success(request, "Insights refreshed.")
    return redirect("dashboard:insights")
